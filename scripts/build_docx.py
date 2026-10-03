# -*- coding: utf-8 -*-
"""把康奈尔错题集 HTML 转成 .docx，保留康奈尔笔记的版面结构（真实表格）。

Usage: python build_docx.py <in.html> <out.docx>
"""
import io
import os
import sys
from bs4 import BeautifulSoup, NavigableString
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NL = chr(10)
LINE = 1.05
FONT = '微软雅黑'
MONO = 'Consolas'


def norm(s):
    return ' '.join(s.split())


def set_run(r, size=10, bold=False, mono=False, color=None):
    r.font.size = Pt(size)
    r.bold = bold
    name = MONO if mono else FONT
    r.font.name = name
    rPr = r._element.get_or_add_rPr()
    rf = rPr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rPr.insert(0, rf)
    rf.set(qn('w:ascii'), name)
    rf.set(qn('w:hAnsi'), name)
    rf.set(qn('w:eastAsia'), name)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return r


def para(container, size=10, space_after=2, space_before=0, line=None, align=None, indent=None):
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = LINE if line is None else line
    if align is not None:
        p.alignment = align
    if indent is not None:
        pf.left_indent = Cm(indent)
    return p


def para_border_bottom(p, sz=6):
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(sz))
    bottom.set(qn('w:space'), '2')
    bottom.set(qn('w:color'), '000000')
    b.append(bottom)
    pPr.append(b)


def table_borders(table, sz=6):
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), str(sz))
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), '000000')
        borders.append(el)
    tblPr.append(borders)


def cell_width(cell, cm):
    cell.width = Cm(cm)
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn('w:tcW'))
    if tcW is None:
        tcW = OxmlElement('w:tcW')
        tcPr.append(tcW)
    tcW.set(qn('w:w'), str(int(cm * 567)))
    tcW.set(qn('w:type'), 'dxa')


def cell_margins(table, top=60, bottom=60, left=90, right=90):
    tblPr = table._tbl.tblPr
    mar = OxmlElement('w:tblCellMar')
    for name, val in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        el = OxmlElement('w:' + name)
        el.set(qn('w:w'), str(val))
        el.set(qn('w:type'), 'dxa')
        mar.append(el)
    tblPr.append(mar)


def one_cell_box(container, width=17.4):
    t = container.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(t)
    cell_margins(t, top=70, bottom=70, left=140, right=140)
    c = t.cell(0, 0)
    c.paragraphs[0]._element.getparent().remove(c.paragraphs[0]._element)
    cell_width(c, width)
    return t, c


def iter_runs(node, bold=False, mono=False):
    for child in node.children:
        if isinstance(child, NavigableString):
            t = norm(str(child))
            if t:
                yield (t, bold, mono)
        elif child.name == 'br':
            yield (NL, bold, mono)
        else:
            b = bold or child.name in ('strong', 'b')
            m = mono or child.name in ('code', 'tt')
            for item in iter_runs(child, b, m):
                yield item


def add_inline(p, node, size=10, base_bold=False, color=None):
    merged = []
    for t, b, m in iter_runs(node, base_bold):
        if merged and merged[-1][1] == b and merged[-1][2] == m:
            merged[-1][0] += t
        else:
            merged.append([t, b, m])
    while merged and merged[0][0].strip() == '':
        merged.pop(0)
    while merged and merged[-1][0].strip() == '':
        merged.pop()
    if merged:
        merged[0][0] = merged[0][0].lstrip()
        merged[-1][0] = merged[-1][0].rstrip()
    for t, b, m in merged:
        if t == '':
            continue
        if NL in t:
            parts = t.split(NL)
            for i, part in enumerate(parts):
                if i:
                    r = p.add_run()
                    set_run(r, size=size, bold=b, mono=m, color=color)
                    r.add_break()
                if part:
                    set_run(p.add_run(part), size=size, bold=b, mono=m, color=color)
        else:
            set_run(p.add_run(t), size=size, bold=b, mono=m, color=color)
    return p


def txt(node):
    return norm(node.get_text())


def build_cover(doc, page):
    p = para(doc, size=22, space_after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_run(p.add_run(txt(page.find('h1'))), size=22, bold=True)
    p = para(doc, size=12, space_after=22, align=WD_ALIGN_PARAGRAPH.CENTER)
    sub = page.find('p', class_='subtitle')
    if sub:
        set_run(p.add_run(txt(sub)), size=12, color='555555')
    src = page.find('table', class_='meta-table')
    t = doc.add_table(rows=0, cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(t)
    cell_margins(t, top=80, bottom=80, left=110, right=110)
    for tr in src.find_all('tr'):
        cells = tr.find_all(['td', 'th'])
        row = t.add_row()
        idx = 0
        for td in cells:
            span = int(td.get('colspan', 1))
            cell = row.cells[idx]
            if span > 1:
                cell = cell.merge(row.cells[idx + span - 1])
            bold = idx in (0, 2)
            cp = cell.paragraphs[0]
            cp.paragraph_format.space_after = Pt(0)
            add_inline(cp, td, size=10, base_bold=bold)
            idx += span
    para(doc, size=6, space_after=10)
    box = page.find('div', class_='how-to-use')
    if not box:
        return
    t, c = one_cell_box(doc)
    p = para(c, size=11, space_after=4)
    set_run(p.add_run(txt(box.find('h3'))), size=11, bold=True)
    para_border_bottom(p)
    for i, li in enumerate(box.find_all('li'), 1):
        p = para(c, size=10, space_after=3, indent=0.6)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        set_run(p.add_run('%d. ' % i), size=10)
        add_inline(p, li, size=10)
    for ptag in box.find_all('p'):
        p = para(c, size=9, space_after=0, space_before=6)
        add_inline(p, ptag, size=9, color='333333')


def build_error_table(doc, src):
    heads = src.find_all('th')
    rows = src.find('tbody').find_all('tr')
    t = doc.add_table(rows=1, cols=len(heads))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(t)
    cell_margins(t)
    for i, th in enumerate(heads):
        c = t.rows[0].cells[i]
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(p, th, size=9.5, base_bold=True)
    for tr in rows:
        row = t.add_row()
        for i, td in enumerate(tr.find_all('td')):
            c = row.cells[i]
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if 'c' in (td.get('class') or []):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, td, size=9)
    return t


def build_box(doc, box, title_size=11, body_size=9):
    t, c = one_cell_box(doc)
    h3 = box.find('h3')
    if h3:
        p = para(c, size=title_size, space_after=4)
        set_run(p.add_run(txt(h3)), size=title_size, bold=True)
        para_border_bottom(p)
    for child in box.children:
        if isinstance(child, NavigableString):
            continue
        if child.name == 'h3':
            continue
        if child.name in ('ol', 'ul'):
            for i, li in enumerate(child.find_all('li', recursive=False), 1):
                p = para(c, size=body_size, space_after=3, indent=0.7)
                p.paragraph_format.first_line_indent = Cm(-0.7)
                set_run(p.add_run('%d. ' % i), size=body_size)
                add_inline(p, li, size=body_size)
        elif child.name in ('p', 'div'):
            p = para(c, size=body_size, space_after=3)
            add_inline(p, child, size=body_size)
    return t


def build_misc(doc, page):
    global LINE
    LINE = 1.05
    title = page.find(class_='page-title')
    p = para(doc, size=15, space_after=6, space_before=0)
    p.paragraph_format.page_break_before = True
    if title:
        set_run(p.add_run(txt(title)), size=15, bold=True)
    para_border_bottom(p, sz=12)
    for child in page.children:
        if isinstance(child, NavigableString):
            continue
        cls = child.get('class') or []
        if child.name == 'div' and 'page-title' in cls:
            continue
        if child.name == 'table':
            build_error_table(doc, child)
            para(doc, size=6, space_after=6)
        elif child.name == 'div':
            build_box(doc, child)
            para(doc, size=6, space_after=6)


def build_question(doc, page):
    global LINE
    LINE = 1.0
    blk = page.find('div', class_='question-block')
    head = blk.find('div', class_='q-header')
    p = para(doc, size=12, space_after=5)
    p.paragraph_format.page_break_before = True
    p.paragraph_format.tab_stops.add_tab_stop(Cm(17.4), WD_TAB_ALIGNMENT.RIGHT)
    set_run(p.add_run(txt(head.find('h2'))), size=12, bold=True)
    meta = head.find('span', class_='q-meta')
    if meta:
        set_run(p.add_run(chr(9)), size=12)
        set_run(p.add_run(txt(meta)), size=8.5, color='555555')
    para_border_bottom(p, sz=12)
    grid = blk.find('div', class_='cornell-grid')
    cue = grid.find('div', class_='cue-column')
    main = grid.find('div', class_='main-column')
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    table_borders(t)
    cell_margins(t, top=30, bottom=30)
    lc, rc = t.rows[0].cells
    cell_width(lc, 4.35)
    cell_width(rc, 13.05)
    lc.paragraphs[0]._element.getparent().remove(lc.paragraphs[0]._element)
    zp = para(lc, size=8, space_after=4)
    set_run(zp.add_run(txt(cue.find(class_='zone-label'))), size=8, bold=True)
    para_border_bottom(zp)
    for f in cue.find_all('div', class_='field'):
        name = f.find(class_='field-name')
        val = f.find(class_='field-value')
        p = para(lc, size=8.5, space_after=3)
        set_run(p.add_run(txt(name)), size=8.5, bold=True)
        if val:
            p.add_run().add_break()
            add_inline(p, val, size=8.5)
    rc.paragraphs[0]._element.getparent().remove(rc.paragraphs[0]._element)
    zp = para(rc, size=8, space_after=4)
    set_run(zp.add_run(txt(main.find(class_='zone-label'))), size=8, bold=True)
    para_border_bottom(zp)
    for sec in main.find_all('div', class_='section'):
        st = sec.find('span', class_='section-title')
        if st:
            p = para(rc, size=8.5, space_after=1, space_before=2)
            set_run(p.add_run(txt(st)), size=8.5, bold=True)
        for d in sec.find_all('div', recursive=False):
            p = para(rc, size=8.5, space_after=1)
            add_inline(p, d, size=8.5)
        if sec.find('hr'):
            p = para(rc, size=4, space_after=1)
            para_border_bottom(p, sz=4)
    para(doc, size=4, space_after=2)
    summ = blk.find('div', class_='summary-column')
    t, c = one_cell_box(doc)
    zp = para(c, size=8, space_after=4)
    set_run(zp.add_run(txt(summ.find(class_='zone-label'))), size=8, bold=True)
    para_border_bottom(zp)
    for f in summ.find_all('div', class_='field'):
        p = para(c, size=8.5, space_after=2)
        add_inline(p, f, size=8.5)


def main():
    if len(sys.argv) < 3:
        print("usage: python build_docx.py <in.html> <out.docx>")
        sys.exit(1)
    src = os.path.abspath(sys.argv[1])
    out = os.path.abspath(sys.argv[2])
    html = io.open(src, encoding='utf-8').read()
    soup = BeautifulSoup(html, 'lxml')
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.6)
    sec.bottom_margin = Cm(1.6)
    sec.left_margin = Cm(1.8)
    sec.right_margin = Cm(1.8)
    st = doc.styles['Normal']
    st.font.name = FONT
    st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    st.paragraph_format.space_after = Pt(2)
    st.paragraph_format.line_spacing = 1.05
    pages = soup.select('div.page')
    nq = 0
    for page in pages:
        cls = page.get('class') or []
        if 'cover' in cls:
            build_cover(doc, page)
        elif page.find('div', class_='question-block'):
            nq += 1
            build_question(doc, page)
        else:
            build_misc(doc, page)
    doc.save(out)
    print('DOCX:', out, os.path.getsize(out), '| question pages:', nq, '| total pages:', len(pages))


if __name__ == '__main__':
    main()

